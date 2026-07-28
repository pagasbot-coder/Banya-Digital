#!/usr/bin/env python3
"""Travel+ Natural interview deck v1.14.3 — copy expand, structured blocks."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import os
import shutil

W = Inches(13.333)
H = Inches(7.5)

BG = RGBColor(0xF5, 0xF3, 0xEF)
ACCENT = RGBColor(0x2A, 0x55, 0x4A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD4, 0xCE, 0xC4)

# Font stack #4: Manrope (display) + Golos Text (body)
FONT_DISPLAY = "Manrope"
FONT_BODY = "Golos Text"
FOOTER_TEXT = "Natural · пример презентации"

MARGIN_X = Inches(0.5)
CONTENT_TOP = Inches(1.05)
CARD_H = Inches(5.55)
GAP = Inches(0.28)
CARD_W = (W - 2 * MARGIN_X - GAP) / 2
STRIP_H = Inches(0.09)
CARD_PAD_X = Inches(0.38)
CARD_PAD_TOP = Inches(0.42)
CARD_PAD_BOTTOM = Inches(0.3)
TF_MARGIN_X = Pt(4)
TF_MARGIN_TOP = Pt(2)
TF_MARGIN_BOTTOM = Pt(2)


def pick_font(size=16, bold=False):
    if bold or size >= 19:
        return FONT_DISPLAY
    return FONT_BODY


def set_run(run, size=16, bold=False, color=INK, name=None):
    name = name or pick_font(size, bold)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("a:rFonts"))
    if rFonts is None:
        rFonts = etree.SubElement(rPr, qn("a:rFonts"))
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(attr, name)


def configure_text_frame(tf, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = TF_MARGIN_X
    tf.margin_right = TF_MARGIN_X
    tf.margin_top = TF_MARGIN_TOP
    tf.margin_bottom = TF_MARGIN_BOTTOM


def add_bg(slide):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shp.fill.solid()
    shp.fill.fore_color.rgb = BG
    shp.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.1), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_footer(slide, page, total=12):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN_X, Inches(6.85), W - 2 * MARGIN_X, Inches(0.015)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    box = slide.shapes.add_textbox(MARGIN_X, Inches(6.95), Inches(10.5), Inches(0.35))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = FOOTER_TEXT
    set_run(r, 11, False, MUTED, FONT_BODY)
    box2 = slide.shapes.add_textbox(Inches(11.5), Inches(6.95), Inches(1.4), Inches(0.35))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = f"{page} / {total}"
    set_run(r2, 11, False, MUTED, FONT_BODY)


def add_title(slide, text):
    box = slide.shapes.add_textbox(MARGIN_X, Inches(0.28), Inches(12.3), Inches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run(r, 26, True, ACCENT)
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN_X, Inches(0.88), Inches(2.2), Inches(0.04)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = ACCENT
    rule.line.fill.background()


def add_card(slide, left, top, width, height):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = LINE
    shp.line.width = Pt(1)
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, STRIP_H)
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()
    return shp


def write_block(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    configure_text_frame(tf)
    tf.clear()
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = item.get("align", align)
        p.space_before = Pt(item.get("before", 0))
        p.space_after = Pt(item.get("after", 8))
        p.line_spacing = item.get("line_spacing", 1.18)
        r = p.add_run()
        r.text = item["text"]
        size = item.get("size", 16)
        bold = item.get("bold", False)
        set_run(r, size, bold, item.get("color", INK), item.get("font", pick_font(size, bold)))
    return box


def card_header(text):
    return {
        "text": text,
        "size": 15,
        "bold": True,
        "color": ACCENT,
        "after": 12,
        "before": 0,
        "font": FONT_DISPLAY,
    }


def body(text, **kw):
    d = {"text": text, "size": 16, "bold": False, "color": INK, "after": 8}
    d.update(kw)
    return d


def muted(text, **kw):
    d = {"text": text, "size": 14, "bold": False, "color": MUTED, "after": 6}
    d.update(kw)
    return d


def strong(text, **kw):
    d = {"text": text, "size": 16, "bold": True, "color": INK, "after": 8}
    d.update(kw)
    return d


def dense(text, **kw):
    d = {"text": text, "size": 15, "bold": False, "color": INK, "after": 6}
    d.update(kw)
    return d


def footnote(text, **kw):
    d = {"text": text, "size": 12, "bold": False, "color": MUTED, "before": 18, "after": 0}
    d.update(kw)
    return d


def labeled(title, text, **kw):
    gap = kw.pop("gap", 10)
    return [
        strong(title, size=15, after=2),
        dense(text, after=gap, **{k: v for k, v in kw.items() if k != "gap"}),
    ]



def card_text_area(left, top, width, height):
    """Text box inset: below accent strip, with horizontal and bottom padding."""
    tx = left + CARD_PAD_X
    ty = top + STRIP_H + CARD_PAD_TOP
    tw = width - 2 * CARD_PAD_X
    th = height - STRIP_H - CARD_PAD_TOP - CARD_PAD_BOTTOM
    return tx, ty, tw, th


def two_cards(slide):
    left = MARGIN_X
    right = MARGIN_X + CARD_W + GAP
    add_card(slide, left, CONTENT_TOP, CARD_W, CARD_H)
    add_card(slide, right, CONTENT_TOP, CARD_W, CARD_H)
    return card_text_area(left, CONTENT_TOP, CARD_W, CARD_H), card_text_area(
        right, CONTENT_TOP, CARD_W, CARD_H
    )


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    add_bg(s)
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(5.1), H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = ACCENT
    panel.line.fill.background()
    write_block(
        s,
        Inches(0.55),
        Inches(2.15),
        Inches(4.2),
        Inches(3.2),
        [
            {"text": "Павел Биджиев", "size": 30, "bold": True, "color": WHITE, "after": 14},
            {"text": "Продакт / бренд-логика", "size": 17, "color": WHITE, "after": 8},
            {
                "text": "Физические товары · HoReCa amenities",
                "size": 14,
                "color": RGBColor(0xC8, 0xDD, 0xD6),
                "after": 0,
            },
        ],
    )
    write_block(
        s,
        Inches(5.7),
        Inches(2.0),
        Inches(6.9),
        Inches(3.8),
        [
            {"text": "Travel+ Natural", "size": 32, "bold": True, "color": ACCENT, "after": 12},
            {"text": "Своя натуральная линейка:", "size": 19, "color": INK, "after": 4},
            {"text": "от рамки до пилота", "size": 19, "color": INK, "after": 18},
            {
                "text": "Подготовка · АК-Сервис / Travel+ · Кириши",
                "size": 14,
                "color": MUTED,
                "after": 6,
            },
            {"text": "Образец презентации · не оферта", "size": 14, "color": MUTED, "after": 0},
        ],
    )

    # 2 Role
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Как понял роль")
    add_footer(s, 2)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Travel+"),
            *labeled("Поставка на номер", "One-stop: косметика, тапочки, наборы"),
            *labeled(
                "Контур производства",
                "Своё производство и сборка; закупки и производство в Китае",
            ),
            *labeled("Фокус роли", "Товар от идеи до полки и маржи.", gap=0),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Роль здесь"),
            *labeled("Ассортимент", "Новинки и позиция в доме брендов"),
            *labeled("Продукт", "Доказательства до выхода на витрину"),
            *labeled("Связка", "Технолог → продажи → экономика → поставка", gap=0),
        ],
    )

    # 3 Gap
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Дыра на витрине")
    add_footer(s, 3)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Рынок"),
            *labeled(
                "Конкуренты",
                "ЕТС Natural · MEZO · Natura Siberica (чужой премиум в каталоге)",
            ),
            *labeled("Сила Travel+", "Массовый и средний сегмент"),
            *labeled("Дыра", "Своей Natural с документами нет", gap=0),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Ценовая лестка (не оферта)"),
            *labeled("Hotel Line", "12–15 ₽ / 30 мл"),
            *labeled("ЕТС Natural", "18 ₽ / 25 мл"),
            *labeled("Natura Siberica", "30–40 ₽ и выше"),
            *labeled("Место Natural", "Между Hotel Line и Natura Siberica", gap=0),
        ],
    )

    # 4 Portfolio
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Дом линеек")
    add_footer(s, 4)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Карта портфеля"),
            dense("Hotel Line — цена и базовый сегмент"),
            dense("Fleur / Aquatique / La Nuit — дизайн и аромат"),
            strong("Natural — натуральность и документы", size=15),
            dense("Natura Siberica — чужой премиум в каталоге"),
            dense("СТМ с логотипом отеля — отдельный проект"),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Мостик для продаж"),
            dense("Fleur закрывает дизайн и аромат в номере."),
            dense("Natural — запрос на натуральность и полный пакет документов."),
            strong("Fleur и Natural не конкурируют — разные задачи.", size=15, after=10),
            dense(
                "На одном объекте: стандартные номера — Hotel Line или Fleur; "
                "wellness-корпус или этаж с eco-запросом — Natural."
            ),
        ],
    )

    # 5 Promise
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Обещание + StoryBrand")
    add_footer(s, 5)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Обещание"),
            strong("Мягкий уход с понятной натуральностью", size=15),
            dense(
                "Для отеля, которому важен комфорт гостя и спокойная проверка документов.",
                after=10,
            ),
            muted("Доказательства:", after=4),
            dense("• Состав в согласованных границах"),
            dense("• Пакет документов Travel+"),
            dense("• Ощущение в номере после использования"),
            dense("• Один поставщик на весь номер"),
            dense("• Блок «простыми словами» на этикетке"),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("BrandScript"),
            strong("Герой", size=15, after=2),
            dense(
                "Закупщик отеля с запросом eco / wellness: "
                "выглядеть «зеленее» без дыры в бюджете и в бумагах.",
                after=8,
            ),
            strong("Проводник", size=15, after=2),
            dense("Travel+ Natural — своя линейка, документы и one-stop на номер.", after=8),
            strong("План", size=15, after=2),
            dense("1. Согласовать сегмент и объект"),
            dense("2. Образец с этикеткой и пакетом документов"),
            dense("3. Пилот на этаже или в корпусе", after=10),
            muted("Гость после дороги:", after=4),
            strong("Ощущение мягче, чем ожидал.", size=15, after=2),
            dense("Понятная натуральность — без громких обещаний на этикетке."),
        ],
    )

    # 6 DoD
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Готовность продукта перед рынком")
    add_footer(s, 6)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Чеклист: без этого — не в продажи"),
            dense("□ Метод расчёта и доля натуральных ингредиентов"),
            dense("□ Пакет документов как у Hotel Line и Fleur + лист состава Natural"),
            dense("□ Происхождение (origin) зафиксировано письменно"),
            dense("□ Слепой тест запаха и ощущения с гостями"),
            dense("□ Этикетка с блоком «простыми словами»"),
            dense("□ Себестоимость и коридор цены на полке"),
            dense("□ Срок поставки, минимальный заказ, правила брака"),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Как читать список"),
            dense("Каждый пункт — обязательное условие."),
            dense("Пока хотя бы один не закрыт — коммерческое предложение и выход на витрину стоп."),
            dense("На старте без знака COSMOS, если его нет — говорим честно.", after=10),
            strong("Этапы 3–6 в пульте — учебный сценарий,", size=15, after=2),
            dense("а не отчёт завода и не основание для КП."),
        ],
    )

    # 7 Pitch
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Питч закупщику (30 секунд)")
    add_footer(s, 7)
    pitch_w = W - 2 * MARGIN_X
    pitch_h = Inches(4.75)
    add_card(s, MARGIN_X, CONTENT_TOP, pitch_w, pitch_h)
    px, py, pw, ph = card_text_area(MARGIN_X, CONTENT_TOP, pitch_w, pitch_h)
    write_block(
        s,
        px,
        py,
        pw,
        ph,
        [
            strong("Задача закупщика", size=16, after=4),
            dense(
                "Отель хочет выглядеть экологичнее — без риска для гостя "
                "и без пробелов в документах при проверке.",
                after=10,
            ),
            strong("Что предлагаем", size=16, after=4),
            dense("Travel+ Natural — пакет документов как у Hotel Line и Fleur:"),
            dense("декларация или СГР, INCI, протоколы по запросу."),
            dense("Состав с блоком «простыми словами». Один договор на весь номер.", after=10),
            strong("Три шага", size=16, after=4),
            dense("1. Сегмент eco / wellness и объект"),
            dense("2. Образец с этикеткой и документами"),
            dense("3. Пилот на этаже или в корпусе"),
        ],
    )
    note_y = CONTENT_TOP + pitch_h + Inches(0.18)
    write_block(
        s,
        MARGIN_X + CARD_PAD_X,
        note_y,
        W - 2 * MARGIN_X - 2 * CARD_PAD_X,
        Inches(0.55),
        [
            muted(
                "Отрицания (COSMOS, происхождение, цена Hotel Line) — только в блоке возражений.",
                after=0,
            )
        ],
    )

    # 8 Objections
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Поле: возражения и риски")
    add_footer(s, 8)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Возражения и ответы"),
            *labeled(
                "COSMOS в тендере обязателен",
                "Знака нет — Natural в это коммерческое предложение не ставим. "
                "Предлагаем Hotel Line или отказ.",
            ),
            *labeled(
                "«Natural по цене Hotel Line»",
                "Нет. Другая себестоимость и другая позиция на полке.",
            ),
            *labeled(
                "СТМ / логотип отеля",
                "Отдельный проект. Не маскируем под Natural.",
                gap=8,
            ),
            *labeled(
                "«Где произведено?»",
                "Только зафиксированное происхождение. Без догадок и «скорее всего».",
                gap=6,
            ),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Полевая дисциплина"),
            dense("Реестр рисков — один лист в пульте (14 пунктов)."),
            dense("Сегмент не размываем: Natural не уходит в эконом-запрос.", after=10),
            strong("Не продаём Natural вместо Hotel Line по цене базовой линейки.", size=15),
            footnote("Цифры сценария — учебный пример, не отчёт завода."),
        ],
    )

    # 9 Pilot
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Пилот: успех, провал, выход")
    add_footer(s, 9)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Успех (учебный порог)"),
            dense("• Три пилотных клиента"),
            dense("• Повторный заказ минимум у двух из трёх"),
            dense("• Мягкость и ощущение натуральности — не ниже 70% в опросе"),
            strong("• Natural не ушёл в эконом Hotel Line", size=15),
            footnote("Цифры сценария — учебный пример, не отчёт завода."),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Провал и выход"),
            dense("Повторный заказ меньше чем у двух из трёх — или метрики ниже 60%."),
            strong("→ Останавливаем масштабирование", size=15, after=8),
            dense("Диагноз бренд-менеджера — в течение двух недель."),
            dense("Второй пилотный круг снова провален."),
            strong("→ Закрываем линейку", size=15),
            footnote("Пороги — для учебного сценария в пульте, не KPI завода."),
        ],
    )

    # 10 Logistics
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Поставки и экономика (Китай)")
    add_footer(s, 10)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Экономика и цена"),
            dense("Прайс утверждаем только после себестоимости на складе в Киришах."),
            dense("Ориентир по полке:", after=4),
            dense("• выше Hotel Line"),
            dense("• рядом с ЕТС Natural"),
            dense("• ниже Natura Siberica", after=10),
            strong("Не воюем копейкой с ЕТС.", size=15, after=4),
            dense("Выигрываем пакетом документов и ощущением в номере."),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Если поставка из Китая"),
            dense("1. Сравниваем заводы в цене FOB"),
            dense("2. Логистику до Кириши считаем отдельной строкой"),
            dense("3. Полная цена = FOB + фрахт + таможня + довоз + запас", after=10),
            strong("Срок пилота — дата «на складе»,", size=15, after=2),
            dense("а не формулировка «завод отгрузил»."),
            dense("Происхождение не зафиксировано — не называем «российское»."),
        ],
    )

    # 11 What's in pult
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Что уже есть в пульте")
    add_footer(s, 11)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Содержание пульта v1.14"),
            dense("Этапы 0–2: разведка, правила, границы продукта"),
            dense("Этапы 3–6: учебный сценарий «как вести дело»"),
            dense("StoryBrand, чеклист готовности, реестр рисков"),
            dense("Правила пилота: успех, провал, выход"),
            dense("Мостик портфеля Travel+ (Hotel Line, Fleur, Natural, СТМ)"),
            footnote("Статус: образец для поля и собеседования. Не коммерческое предложение."),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Справочники и логистика"),
            dense("Глоссарий бренда и глоссарий ВЭД"),
            dense("Раздел L: путь товара, Incoterms, RFQ, приёмка на складе"),
            dense("Волна A: markdown-реестры в папке brand-pult/"),
            dense("Презентация и скрипт — синхрон с пультом", after=10),
            strong("Версия 2.0 — когда технолог закроет цифры завода.", size=15),
        ],
    )

    # 12 Close
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_title(s, "Итог и следующий шаг")
    add_footer(s, 12)
    L, R = two_cards(s)
    write_block(
        s,
        *L,
        [
            card_header("Собрал"),
            body("ДНК + StoryBrand + DoD"),
            body("Полевая защита"),
            body("Чеклист поставок при Китае (раздел L)", after=14),
            muted("Не прайс завода. Не КП."),
        ],
    )
    write_block(
        s,
        *R,
        [
            card_header("Готов обсудить"),
            body("1) Приоритет блокеров этапа 2 с технологом"),
            body("2) Пилотный сегмент eco / wellness"),
            body("3) Как вести продажи, чтобы Natural не ушёл в эконом Hotel Line", after=14),
            strong("Спасибо. Вопросы."),
        ],
    )

    out = "knowledge-base/presentation-travelplus-interview.pptx"
    prs.save(out)
    dests = [
        "АК/presentation-travelplus-interview.pptx",
        "/opt/cursor/artifacts/presentation-travelplus-interview.pptx",
        "/opt/cursor/artifacts/АК/presentation-travelplus-interview.pptx",
        "/home/ubuntu/АК/presentation-travelplus-interview.pptx",
        "/home/ubuntu/Desktop/АК/presentation-travelplus-interview.pptx",
    ]
    for dest in dests:
        os.makedirs(os.path.dirname(dest) or ".", exist_ok=True)
        shutil.copy2(out, dest)
    print("saved", out, "slides", len(prs.slides))
    return out


if __name__ == "__main__":
    build()
